"""Validate the actual XLSX bytes generated for capper exports."""
from io import BytesIO
import unittest

from openpyxl import load_workbook
from bettracker_excel import export_xlsx


class ExcelExportTests(unittest.TestCase):
    def test_summary_bets_filters_and_discord_links(self):
        headers = [
            "bet_date", "graded_utc", "capper", "sport", "wager_category",
            "platform", "player", "market", "odds_text", "risk_units",
            "result", "net_units", "jump_url", "content",
        ]
        bets = [
            ("2026-10-30", "2026-10-30T22:00:00", "PropKitchen", "NBA", "STRAIGHT",
             "FanDuel", "Andrew Nembhard", "o8.5 Assists", "-138", 1.0,
             "win", 0.7246, "https://discord.com/channels/1/2/3", "1u | NBA"),
            ("2026-10-30", "2026-10-30T23:00:00", "PropKitchen", "NBA", "DFS",
             "PrizePicks", "2-leg slip", "Flex 5/6", "2x", .5,
             "win", .5, "https://discord.com/channels/1/2/4", "0.5u | DFS"),
            ("2026-10-30", "2026-10-30T23:10:00", "clipset", "NFL", "DFS",
             "Chalkboard", "", "Slip", "2x", .5,
             "void", 0, "", "=HYPERLINK(\"evil\")"),
        ]
        data = export_xlsx(headers, bets, "October 2026")
        self.assertGreater(len(data), 1000)
        workbook = load_workbook(BytesIO(data), data_only=False)
        self.assertEqual(workbook.sheetnames, ["Summary", "Bets"])
        summary, sheet = workbook["Summary"], workbook["Bets"]
        self.assertEqual(summary["A6"].value, 3)
        self.assertEqual(summary["C6"].value, 2)
        self.assertEqual(summary["I6"].value, 1)
        self.assertEqual(summary["A8"].value, 1.2246)
        self.assertEqual(sheet.max_row, 4)
        self.assertEqual(sheet["A1"].value, "Date")
        self.assertEqual(sheet["L2"].value, "Open Bet")
        self.assertEqual(sheet["L2"].hyperlink.target, "https://discord.com/channels/1/2/3")
        self.assertIsNone(sheet["L4"].hyperlink)
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, "A1:L4")
        self.assertFalse(sheet["F4"].data_type == "f")

    def test_empty_export_still_has_summary_and_headers(self):
        data = export_xlsx(["result"], [], "None")
        book = load_workbook(BytesIO(data))
        self.assertEqual(book["Summary"]["A6"].value, 0)
        self.assertEqual(book["Bets"]["A1"].value, "Date")


if __name__ == "__main__":
    unittest.main()
