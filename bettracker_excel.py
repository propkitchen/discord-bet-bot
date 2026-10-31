"""Polished BetTracker Excel export; no Discord or database dependencies."""
from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from io import BytesIO
import re
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


NAVY = "142039"
BLUE = "2563EB"
WHITE = "FFFFFF"
GREEN = "15803D"
RED = "B91C1C"
PALE = "EFF6FF"
GRAY = "475569"


def safe_text(value: object, limit: int = 360) -> str:
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value or "")).strip()
    text = re.sub(r"<@!?&?\d+>", "", text)
    text = re.sub(r"https?://\S+", "", text)
    text = re.sub(r"\s+", " ", text).strip()
    if text.startswith("="):
        text = "'" + text
    return text[:limit]


def to_float(value: object) -> float:
    try:
        return float(value or 0)
    except (ValueError, TypeError):
        return 0.0


def date_cell(value: object) -> date | str:
    text = str(value or "").strip()
    try:
        return datetime.strptime(text[:10], "%Y-%m-%d").date()
    except ValueError:
        return safe_text(text, 20)


def category_label(value: object) -> str:
    mapping = {"STRAIGHT": "Straight", "SPORTSBOOK_PARLAY": "Sportsbook Parlay", "DFS": "DFS Slip"}
    return mapping.get(str(value or "").upper(), safe_text(value, 40) or "Straight")


def export_xlsx(headers: Sequence[str], rows: Sequence[Sequence[object]], period_label: str) -> bytes:
    """Return an Excel workbook built only from the current server's scoped query."""
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    bets = wb.create_sheet("Bets")
    wb.calculation.fullCalcOnLoad = True

    summary.merge_cells("A1:L2")
    top = summary["A1"]
    top.value = "BETTRACKER  |  PERFORMANCE EXPORT"
    top.font = Font(name="Aptos Display", size=19, bold=True, color=WHITE)
    top.fill = PatternFill("solid", fgColor=NAVY)
    top.alignment = Alignment(vertical="center", indent=1)
    summary.row_dimensions[1].height = 22
    summary.row_dimensions[2].height = 20
    summary.merge_cells("A3:L3")
    summary["A3"] = f"Period: {safe_text(period_label, 100)}  •  Official graded wagers"
    summary["A3"].font = Font(name="Aptos", italic=True, color=GRAY, size=11)
    summary.row_dimensions[3].height = 24
    for col in range(1, 13):
        summary.column_dimensions[get_column_letter(col)].width = 15

    data = [dict(zip(headers, row)) for row in rows]
    rollup = defaultdict(lambda: {"bets": 0, "wins": 0, "losses": 0, "pushes": 0,
                                  "voids": 0, "risk": 0.0, "net": 0.0})
    for row in data:
        capper = safe_text(row.get("capper"), 80) or "Unknown"
        group = rollup[capper]
        group["bets"] += 1
        result = str(row.get("result") or "").lower()
        if result == "win":
            group["wins"] += 1
        elif result == "loss":
            group["losses"] += 1
        elif result == "push":
            group["pushes"] += 1
        elif result == "void":
            group["voids"] += 1
        if result != "void":
            group["risk"] += to_float(row.get("risk_units"))
        group["net"] += to_float(row.get("net_units"))

    total = {k: sum(x[k] for x in rollup.values()) for k in
             ("bets", "wins", "losses", "pushes", "voids", "risk", "net")}
    graded = total["wins"] + total["losses"]
    metrics = [
        ("TOTAL BETS", total["bets"], "A5"),
        ("WINS", total["wins"], "C5"),
        ("LOSSES", total["losses"], "E5"),
        ("PUSHES", total["pushes"], "G5"),
        ("VOIDS", total["voids"], "I5"),
        ("WIN RATE", total["wins"] / graded if graded else 0, "K5"),
        ("NET UNITS", total["net"], "A7"),
        ("ROI", total["net"] / total["risk"] if total["risk"] else 0, "C7"),
        ("RISKED UNITS", total["risk"], "E7"),
    ]
    for label, value, addr in metrics:
        cell = summary[addr]
        cell.value = label
        cell.fill = PatternFill("solid", fgColor=PALE)
        cell.font = Font(name="Aptos", bold=True, color=GRAY, size=10)
        cell.alignment = Alignment(indent=1)
        below = summary.cell(row=cell.row + 1, column=cell.column)
        below.value = value
        below.font = Font(name="Aptos Display", bold=True, size=16,
                          color=GREEN if label == "NET UNITS" and value >= 0 else NAVY)
        below.alignment = Alignment(indent=1)
        if label in {"WIN RATE", "ROI"}:
            below.number_format = "0.0%"
        elif label in {"NET UNITS", "RISKED UNITS"}:
            below.number_format = '+0.00"u";-0.00"u";0.00"u"'
    summary.row_dimensions[5].height = 23
    summary.row_dimensions[6].height = 27
    summary.row_dimensions[7].height = 23
    summary.row_dimensions[8].height = 27

    summary["A11"] = "CAPPER BREAKDOWN"
    summary["A11"].font = Font(name="Aptos Display", bold=True, size=13, color=NAVY)
    rollup_headers = ["Capper", "Bets", "Wins", "Losses", "Pushes", "Voids", "Net Units", "ROI"]
    for col, label in enumerate(rollup_headers, 1):
        cell = summary.cell(row=13, column=col, value=label)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row_id, (capper, numbers) in enumerate(
        sorted(rollup.items(), key=lambda x: (-x[1]["net"], x[0].lower())), 14
    ):
        risk = numbers["risk"]
        values = [capper, numbers["bets"], numbers["wins"], numbers["losses"],
                  numbers["pushes"], numbers["voids"], numbers["net"],
                  numbers["net"] / risk if risk else 0]
        for col_id, value in enumerate(values, 1):
            cell = summary.cell(row=row_id, column=col_id, value=value)
            if row_id % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F8FAFC")
            if col_id == 7:
                cell.number_format = '+0.00"u";-0.00"u";0.00"u"'
            if col_id == 8:
                cell.number_format = "0.0%"
        summary.row_dimensions[row_id].height = 20
    summary.column_dimensions["A"].width = 25
    summary.column_dimensions["G"].width = 18
    summary.freeze_panes = "B14"

    fields = [
        "Date", "Capper", "Sport", "Wager Type", "Platform", "Player / Bet",
        "Market", "Odds / Multiplier", "Risk (u)", "Result", "Net (u)",
        "Discord Link",
    ]
    for idx, label in enumerate(fields, 1):
        cell = bets.cell(1, idx, label)
        cell.font = Font(name="Aptos", bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    bets.row_dimensions[1].height = 32
    bets.freeze_panes = "A2"
    widths = [14, 20, 12, 20, 18, 38, 32, 20, 13, 12, 14, 17]
    for idx, width in enumerate(widths, 1):
        bets.column_dimensions[get_column_letter(idx)].width = width

    for idx, row in enumerate(data, 2):
        url = str(row.get("jump_url") or "").strip()
        bet_text = safe_text(row.get("player"), 120)
        if not bet_text:
            bet_text = safe_text(row.get("market") or row.get("content"), 280)
        values = [
            date_cell(row.get("bet_date") or row.get("graded_utc")),
            safe_text(row.get("capper"), 80),
            safe_text(row.get("sport"), 20),
            category_label(row.get("normalized_wager_category") or row.get("wager_category")),
            safe_text(row.get("platform") or row.get("sportsbook"), 50),
            bet_text,
            safe_text(row.get("market"), 250),
            safe_text(row.get("odds_text"), 40),
            to_float(row.get("risk_units")),
            safe_text(row.get("result"), 20).title(),
            to_float(row.get("net_units")),
            "Open Bet" if url.startswith("https://discord.com/channels/") else "",
        ]
        for col, value in enumerate(values, 1):
            cell = bets.cell(idx, col, value)
            cell.font = Font(name="Aptos", size=10, color="1E293B")
            cell.alignment = Alignment(vertical="top", wrap_text=col in (6, 7))
            if idx % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F8FAFC")
        bets.cell(idx, 1).number_format = "mmm d, yyyy"
        bets.cell(idx, 9).number_format = '0.00"u"'
        bets.cell(idx, 11).number_format = '+0.00"u";-0.00"u";0.00"u"'
        if values[11]:
            jump = bets.cell(idx, 12)
            jump.hyperlink = url
            jump.font = Font(name="Aptos", color=BLUE, underline="single")
        net_cell = bets.cell(idx, 11)
        if values[10] < 0:
            net_cell.font = Font(name="Aptos", color=RED, bold=True)
        elif values[10] > 0:
            net_cell.font = Font(name="Aptos", color=GREEN, bold=True)
        bets.row_dimensions[idx].height = 34
    bets.auto_filter.ref = f"A1:L{max(2, len(data) + 1)}"
    bets.sheet_view.showGridLines = False
    summary.sheet_view.showGridLines = False
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
